
import os
import json
import re
from pathlib import Path

import pymupdf
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from docx import Document

# Load .env from the project root.
BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")


class Candidate(BaseModel):
    name: str
    match_score: float = Field(ge=0, le=100)
    matched_skills: list[str] = Field(default_factory=list)
    missing_skills: list[str] = Field(default_factory=list)
    detected_skills: list[str] = Field(default_factory=list)
    additional_skills: list[str] = Field(default_factory=list)
    strengths: list[str] = Field(default_factory=list)
    overall_assessment: str


# Canonical skill names and their common aliases.
SKILL_ALIASES = {
    # Programming languages
    "Python": ["python", "python programming"],
    "Java": ["java"],
    "JavaScript": ["javascript", "js"],
    "TypeScript": ["typescript", "ts"],
    "C": ["c programming"],
    "C++": ["c++", "cpp"],
    "C#": ["c#", "c sharp"],
    "Go": ["golang", "go language"],
    "Rust": ["rust language"],
    "PHP": ["php"],
    "Ruby": ["ruby"],
    "Swift": ["swift programming"],
    "Kotlin": ["kotlin"],
    "R": ["r programming", "r language"],
    "Scala": ["scala"],
    "Dart": ["dart programming"],
    "MATLAB": ["matlab"],
    "Bash": ["bash", "shell scripting"],
    "PowerShell": ["powershell"],

    # Frontend
    "HTML": ["html", "html5"],
    "CSS": ["css", "css3"],
    "Sass": ["sass", "scss"],
    "Tailwind CSS": ["tailwind css", "tailwind"],
    "Bootstrap": ["bootstrap"],
    "React": ["react", "react.js", "reactjs"],
    "Next.js": ["next.js", "nextjs"],
    "Angular": ["angular"],
    "Vue.js": ["vue.js", "vuejs"],
    "Vite": ["vite"],
    "Redux": ["redux", "redux toolkit"],
    "jQuery": ["jquery"],
    "Responsive Design": ["responsive design", "mobile-first design"],

    # Backend
    "Node.js": ["node.js", "nodejs"],
    "Express.js": ["express.js", "expressjs"],
    "Django": ["django"],
    "Django REST Framework": ["django rest framework", "drf"],
    "Flask": ["flask"],
    "FastAPI": ["fastapi", "fast api"],
    "Spring Boot": ["spring boot"],
    "Spring Framework": ["spring framework"],
    "ASP.NET Core": ["asp.net core", "asp.net"],
    "Laravel": ["laravel"],
    "Symfony": ["symfony"],
    "Ruby on Rails": ["ruby on rails", "rails framework"],
    "NestJS": ["nestjs", "nest.js"],
    "Microservices": ["microservices", "microservice architecture"],

    # APIs and architecture
    "REST API": ["rest api", "rest apis", "restful api", "restful apis"],
    "GraphQL": ["graphql"],
    "WebSockets": ["websocket", "websockets"],
    "API Development": ["api development", "api design"],
    "API Integration": ["api integration", "third-party api integration"],
    "JSON": ["json"],
    "XML": ["xml"],
    "OAuth": ["oauth", "oauth 2.0"],
    "JWT": ["jwt", "json web token"],
    "OpenAPI": ["openapi", "swagger"],
    "API Documentation": ["api documentation"],
    "System Design": ["system design"],
    "Object-Oriented Programming": ["object-oriented programming", "oop"],
    "Data Structures": ["data structures"],
    "Algorithms": ["algorithms"],
    "Design Patterns": ["design patterns"],
    "Caching": ["caching", "cache design"],

    # Databases
    "SQL": ["sql", "structured query language"],
    "PostgreSQL": ["postgresql", "postgres"],
    "MySQL": ["mysql"],
    "SQLite": ["sqlite"],
    "Microsoft SQL Server": ["sql server", "microsoft sql server"],
    "Oracle Database": ["oracle database", "oracle db"],
    "MongoDB": ["mongodb", "mongo db"],
    "Redis": ["redis"],
    "Elasticsearch": ["elasticsearch", "elastic search"],
    "DynamoDB": ["dynamodb"],
    "Firebase": ["firebase"],
    "Supabase": ["supabase"],
    "Cassandra": ["cassandra"],
    "MariaDB": ["mariadb"],
    "Database Design": ["database design", "schema design"],
    "Query Optimization": ["query optimization", "sql optimization"],
    "Database Indexing": ["database indexing", "index optimization"],

    # Cloud and DevOps
    "AWS": ["aws", "amazon web services"],
    "Microsoft Azure": ["microsoft azure", "azure cloud"],
    "Google Cloud": ["google cloud", "gcp"],
    "Docker": ["docker", "containerization"],
    "Kubernetes": ["kubernetes", "k8s"],
    "Jenkins": ["jenkins"],
    "GitHub Actions": ["github actions"],
    "GitLab CI/CD": ["gitlab ci", "gitlab ci/cd"],
    "CI/CD": ["ci/cd", "continuous integration", "continuous deployment"],
    "Terraform": ["terraform"],
    "Ansible": ["ansible"],
    "Linux": ["linux"],
    "Nginx": ["nginx"],
    "Apache": ["apache server"],
    "Render": ["render.com", "render deployment"],
    "Railway": ["railway deployment"],
    "Vercel": ["vercel"],
    "Netlify": ["netlify"],
    "Monitoring": ["monitoring", "application monitoring"],
    "Logging": ["logging", "application logging"],

    # Version control and tools
    "Git": ["git", "git version control"],
    "GitHub": ["github"],
    "GitLab": ["gitlab"],
    "Bitbucket": ["bitbucket"],
    "Postman": ["postman"],
    "Insomnia": ["insomnia api client"],
    "VS Code": ["visual studio code", "vs code"],
    "PyCharm": ["pycharm"],
    "Jira": ["jira"],
    "Confluence": ["confluence"],

    # Testing and quality
    "Unit Testing": ["unit testing", "unit tests"],
    "Integration Testing": ["integration testing"],
    "Pytest": ["pytest"],
    "Unittest": ["python unittest"],
    "Jest": ["jest testing", "jest"],
    "React Testing Library": ["react testing library"],
    "Selenium": ["selenium"],
    "Cypress": ["cypress testing"],
    "Playwright": ["playwright testing"],
    "Test Automation": ["test automation", "automated testing"],
    "Debugging": ["debugging", "troubleshooting"],
    "Code Review": ["code review", "peer review"],
    "Agile": ["agile methodology", "agile"],
    "Scrum": ["scrum"],

    # Data and analytics
    "NumPy": ["numpy"],
    "Pandas": ["pandas"],
    "Matplotlib": ["matplotlib"],
    "Seaborn": ["seaborn"],
    "Plotly": ["plotly"],
    "Excel": ["microsoft excel", "excel"],
    "Power BI": ["power bi"],
    "Tableau": ["tableau"],
    "ETL": ["etl", "extract transform load"],
    "Data Cleaning": ["data cleaning", "data cleansing"],
    "Data Visualization": ["data visualization"],
    "Statistics": ["statistics", "statistical analysis"],
    "Exploratory Data Analysis": ["exploratory data analysis", "eda"],

    # AI, ML, NLP
    "Machine Learning": ["machine learning", "ml"],
    "Deep Learning": ["deep learning"],
    "Artificial Intelligence": ["artificial intelligence", "ai"],
    "Natural Language Processing": ["natural language processing", "nlp"],
    "Computer Vision": ["computer vision"],
    "Generative AI": ["generative ai", "genai"],
    "Large Language Models": ["large language models", "llms"],
    "Prompt Engineering": ["prompt engineering"],
    "LangChain": ["langchain"],
    "LlamaIndex": ["llamaindex", "llama index"],
    "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"],
    "Scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "Keras": ["keras"],
    "Hugging Face": ["hugging face", "transformers library"],
    "OpenCV": ["opencv"],
    "RAG": ["retrieval augmented generation", "rag"],
    "Vector Databases": ["vector database", "vector databases"],
    "Model Evaluation": ["model evaluation", "model validation"],
    "Feature Engineering": ["feature engineering"],

    # Security
    "Cybersecurity": ["cybersecurity", "cyber security"],
    "Application Security": ["application security"],
    "Authentication": ["authentication"],
    "Authorization": ["authorization"],
    "Encryption": ["encryption"],
    "OWASP": ["owasp"],
    "Secure Coding": ["secure coding"],
}


def contains_phrase(text: str, phrase: str) -> bool:
    """Match a phrase with boundaries to avoid substring false positives."""
    text = text.casefold()
    phrase = phrase.casefold().strip()

    if not phrase:
        return False

    pattern = r"(?<!\w)" + re.escape(phrase) + r"(?!\w)"
    return re.search(pattern, text) is not None


def extract_resume_text(file_path: Path) -> str:
    """Extract text from PDF or DOCX resumes."""
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        with pymupdf.open(file_path) as document:
            return "\n".join(page.get_text() for page in document)

    if suffix == ".docx":
        document = Document(file_path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]

        table_text = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]

        return "\n".join(paragraphs + table_text)

    raise ValueError("Unsupported format. Use PDF or DOCX.")


def extract_skills(resume_text: str) -> list[str]:
    """Detect skills and return unique canonical skill names."""
    detected = []

    for skill, aliases in SKILL_ALIASES.items():
        if any(contains_phrase(resume_text, alias) for alias in aliases):
            detected.append(skill)

    return detected


# Broad skill families. These are categories, not synonyms for technologies.
SKILL_GROUPS = {
    "Programming Languages": [
        "Python", "Java", "JavaScript", "TypeScript", "C", "C++", "C#",
        "Go", "Rust", "PHP", "Ruby", "Swift", "Kotlin", "R", "Scala",
        "Dart", "MATLAB", "Bash", "PowerShell",
    ],
    "Frontend Development": [
        "HTML", "CSS", "Sass", "Tailwind CSS", "Bootstrap", "React",
        "Next.js", "Angular", "Vue.js", "Vite", "Redux", "jQuery",
        "Responsive Design",
    ],
    "Backend Development": [
        "Node.js", "Express.js", "Django", "Django REST Framework", "Flask",
        "FastAPI", "Spring Boot", "Spring Framework", "ASP.NET Core",
        "Laravel", "Symfony", "Ruby on Rails", "NestJS", "Microservices",
    ],
    "API Development": [
        "REST API", "GraphQL", "WebSockets", "API Development",
        "API Integration", "JSON", "XML", "OAuth", "JWT", "OpenAPI",
        "API Documentation",
    ],
    "Relational Databases": [
        "PostgreSQL", "MySQL", "SQLite", "Microsoft SQL Server",
        "Oracle Database", "MariaDB", "SQL", "Database Design",
        "Query Optimization", "Database Indexing",
    ],
    "NoSQL and Caching": [
        "MongoDB", "Redis", "Elasticsearch", "DynamoDB", "Firebase",
        "Supabase", "Cassandra", "Caching", "Vector Databases",
    ],
    "Cloud and DevOps": [
        "AWS", "Microsoft Azure", "Google Cloud", "Docker", "Kubernetes",
        "Jenkins", "GitHub Actions", "GitLab CI/CD", "CI/CD", "Terraform",
        "Ansible", "Linux", "Nginx", "Apache", "Render", "Railway",
        "Vercel", "Netlify", "Monitoring", "Logging",
    ],
    "Version Control and Developer Tools": [
        "Git", "GitHub", "GitLab", "Bitbucket", "Postman", "Insomnia",
        "VS Code", "PyCharm", "Jira", "Confluence",
    ],
    "Testing and Quality": [
        "Unit Testing", "Integration Testing", "Pytest", "Unittest", "Jest",
        "React Testing Library", "Selenium", "Cypress", "Playwright",
        "Test Automation", "Debugging", "Code Review", "Agile", "Scrum",
    ],
    "Data and Analytics": [
        "NumPy", "Pandas", "Matplotlib", "Seaborn", "Plotly", "Excel",
        "Power BI", "Tableau", "ETL", "Data Cleaning", "Data Visualization",
        "Statistics", "Exploratory Data Analysis",
    ],
    "AI and Machine Learning": [
        "Machine Learning", "Deep Learning", "Artificial Intelligence",
        "Natural Language Processing", "Computer Vision", "Generative AI",
        "Large Language Models", "Prompt Engineering", "LangChain",
        "LlamaIndex", "TensorFlow", "PyTorch", "Scikit-learn", "Keras",
        "Hugging Face", "OpenCV", "RAG", "Model Evaluation",
        "Feature Engineering",
    ],
    "Security": [
        "Cybersecurity", "Application Security", "Authentication",
        "Authorization", "Encryption", "OWASP", "Secure Coding",
    ],
}


GROUP_ALIASES = {
    "Programming Languages": [
        "programming language", "programming languages",
        "coding language", "coding languages",
    ],
    "Frontend Development": [
        "frontend", "front-end", "frontend development",
        "front end development", "ui development",
    ],
    "Backend Development": [
        "backend", "back-end", "backend development",
        "back end development", "server-side development",
    ],
    "API Development": [
        "api development", "api integration", "apis",
    ],
    "Relational Databases": [
        "relational database", "relational databases", "relational db",
        "sql databases", "database management",
    ],
    "NoSQL and Caching": [
        "nosql", "nosql databases", "no sql databases",
        "caching technologies",
    ],
    "Cloud and DevOps": [
        "cloud computing", "cloud platforms", "devops", "cloud and devops",
    ],
    "Version Control and Developer Tools": [
        "version control", "developer tools",
    ],
    "Testing and Quality": [
        "software testing", "testing frameworks", "quality assurance",
        "qa testing",
    ],
    "Data and Analytics": [
        "data analytics", "data analysis", "business intelligence",
    ],
    "AI and Machine Learning": [
        "artificial intelligence and machine learning", "ai/ml",
        "ai and machine learning",
    ],
    "Security": [
        "application security", "information security", "security practices",
    ],
}


def skill_group_for(skill: str) -> str:
    """Return a skill's display group."""
    for group, skills in SKILL_GROUPS.items():
        if skill in skills:
            return group

    return "Other Skills"


def match_skills(
    candidate_skills: list[str],
    job_description: str,
) -> tuple[list[str], list[str]]:
    """Match broad requirements using related skills, but keep specific
    technology requirements distinct."""

    candidate_set = set(candidate_skills)
    matched = []
    missing = []
    requirements = []

    # Explicitly named skill requirements
    for skill, aliases in SKILL_ALIASES.items():
        if any(contains_phrase(job_description, alias) for alias in aliases):
            requirements.append((skill, "specific"))

    # Broad skill-family requirements
    for group, aliases in GROUP_ALIASES.items():
        if any(contains_phrase(job_description, alias) for alias in aliases):
            requirements.append((group, "group"))

    # SQL is a query language; relational databases can provide evidence
    # of SQL experience when the job asks for SQL generally.
    sql_related = {
        "SQL", "PostgreSQL", "MySQL", "SQLite",
        "Microsoft SQL Server", "Oracle Database", "MariaDB",
    }

    seen = set()

    for label, kind in requirements:
        key = label.casefold()
        if key in seen:
            continue
        seen.add(key)

        if kind == "group":
            related_skills = SKILL_GROUPS.get(label, [])
            is_matched = any(
                skill in candidate_set for skill in related_skills
            )
        elif label == "SQL":
            is_matched = any(
                skill in candidate_set for skill in sql_related
            )
        else:
            # Specific technology requirements must match exactly.
            is_matched = label in candidate_set

        if is_matched:
            matched.append(label)
        else:
            missing.append(label)

    return matched, missing



def calculate_match_score(
    matched_skills: list[str],
    missing_skills: list[str],
) -> float:
    """Calculate the percentage of recognized requirements found."""
    total = len(matched_skills) + len(missing_skills)

    if total == 0:
        return 0.0

    return round(len(matched_skills) / total * 100, 2)



def assess_with_ai(
    resume_text: str,
    job_description: str,
    matched_skills: list[str],
    missing_skills: list[str],
    detected_skills: list[str],
    additional_skills: list[str],
) -> Candidate:
    """Generate an evidence-based assessment using Groq."""
    api_key = os.getenv("GROQ_API_KEY")

    if not api_key:
        raise ValueError("GROQ_API_KEY is missing from .env")

    client = Groq(api_key=api_key)

    response = client.chat.completions.create(
        model="openai/gpt-oss-20b",
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a resume assessment assistant. "
                    "Use only evidence from the supplied resume. "
                    "Do not invent experience, qualifications, or skills. "
                    "Return valid JSON only."
                ),
            },
            {
                "role": "user",
                "content": f"""
Analyze the candidate against the job description.

Return JSON with exactly these fields:
- name: candidate's name, or "Unknown Candidate"
- strengths: a list of 2-5 evidence-based strengths
- overall_assessment: a concise assessment of the candidate's fit

Resume:
{resume_text[:25000]}

Job description:
{job_description[:10000]}

Matched requirements:
{json.dumps(matched_skills)}

Missing requirements:
{json.dumps(missing_skills)}

Detected skills:
{json.dumps(detected_skills)}
""",
            },
        ],
        response_format={"type": "json_object"},
    )

    content = response.choices[0].message.content

    if not content:
        raise ValueError("The AI provider returned an empty response.")

    ai_data = json.loads(content)

    name = ai_data.get("name")
    if not isinstance(name, str) or not name.strip():
        name = "Unknown Candidate"

    strengths = ai_data.get("strengths")
    if not isinstance(strengths, list):
        strengths = []

    strengths = [
        item.strip()
        for item in strengths
        if isinstance(item, str) and item.strip()
    ]

    assessment = ai_data.get("overall_assessment")

    if not isinstance(assessment, str) or not assessment.strip():
        if missing_skills:
            assessment = (
                "The resume demonstrates some relevant skills, but "
                "does not show evidence for all recognized requirements. "
                "Missing requirements include: "
                + ", ".join(missing_skills)
                + "."
            )
        else:
            assessment = (
                "All recognized requirements were detected in the resume. "
                "Review the candidate's experience and project evidence "
                "to confirm practical proficiency."
            )

    return Candidate(
        name=name.strip(),
        match_score=calculate_match_score(
            matched_skills, missing_skills
        ),
        matched_skills=matched_skills,
        missing_skills=missing_skills,
        detected_skills=detected_skills,
        additional_skills=additional_skills,
        strengths=strengths,
        overall_assessment=assessment.strip(),
    )


def screen_resume(file_path: Path, job_description: str) -> Candidate:
    """Run the full screening pipeline for one resume."""
    resume_text = extract_resume_text(file_path)

    if not resume_text.strip():
        raise ValueError(
            "No selectable text found. This PDF may be scanned or image-based."
        )

    candidate_skills = extract_skills(resume_text)

    matched, missing = match_skills(candidate_skills, job_description)

    # Skills not explicitly required by the job are listed separately.
    explicit_required_skills = {
        skill
        for skill, aliases in SKILL_ALIASES.items()
        if any(contains_phrase(job_description, alias) for alias in aliases)
    }

    additional = [
        skill for skill in candidate_skills
        if skill not in explicit_required_skills
    ]

    return assess_with_ai(
        resume_text=resume_text,
        job_description=job_description,
        matched_skills=matched,
        missing_skills=missing,
        detected_skills=candidate_skills,
        additional_skills=additional,
    )


if __name__ == "__main__":
    print(contains_phrase("Experience with SQL and MySQL", "SQL"))
    print(contains_phrase("Used SQL queries", "SQL"))
    print(contains_phrase("No database experience", "SQL"))
